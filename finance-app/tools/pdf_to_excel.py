#!/usr/bin/env python3
"""
pdf_to_excel.py — Convert a PDF statement to Excel, preserving table structure.

Usage:
    python3 pdf_to_excel.py <input.pdf> [output.xlsx]

Dependencies (install once):
    pip install pdfplumber openpyxl

Optional (better table detection for complex PDFs):
    pip install camelot-py[cv] ghostscript
"""

import sys
import os
import argparse
from pathlib import Path


def install_and_import(package, import_name=None):
    import importlib
    name = import_name or package
    try:
        return importlib.import_module(name)
    except ImportError:
        print(f"[INFO] Installing {package}...")
        import subprocess
        subprocess.check_call([sys.executable, "-m", "pip", "install", package, "-q"])
        return importlib.import_module(name)


def extract_tables_pdfplumber(pdf_path):
    """Extract all tables from PDF using pdfplumber. Returns list of (page_num, table_index, rows)."""
    pdfplumber = install_and_import("pdfplumber")
    results = []
    with pdfplumber.open(pdf_path) as pdf:
        print(f"[INFO] PDF has {len(pdf.pages)} page(s)")
        for page_num, page in enumerate(pdf.pages, start=1):
            tables = page.extract_tables()
            if tables:
                for tbl_idx, table in enumerate(tables):
                    # Filter out fully-empty rows
                    clean = [row for row in table if any(cell and str(cell).strip() for cell in row)]
                    if clean:
                        results.append((page_num, tbl_idx + 1, clean))
                        print(f"  Page {page_num}: table {tbl_idx+1} — {len(clean)} rows x {len(clean[0])} cols")
            else:
                # Fall back to raw text lines for non-tabular pages
                text = page.extract_text()
                if text and text.strip():
                    lines = [[line] for line in text.splitlines() if line.strip()]
                    results.append((page_num, 0, lines))
                    print(f"  Page {page_num}: no tables found — captured {len(lines)} text lines")
    return results


def write_excel(tables, output_path):
    """Write extracted tables to Excel. Each page gets its own sheet."""
    openpyxl = install_and_import("openpyxl")
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter

    wb = openpyxl.Workbook()
    wb.remove(wb.active)  # remove default sheet

    HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
    HEADER_FONT = Font(bold=True, color="FFFFFF", size=10)
    ALT_FILL    = PatternFill("solid", fgColor="D6E4F0")
    NORM_FILL   = PatternFill("solid", fgColor="FFFFFF")
    THIN        = Side(style="thin", color="BFBFBF")
    BORDER      = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

    # Group by page
    from collections import defaultdict
    pages = defaultdict(list)
    for page_num, tbl_idx, rows in tables:
        pages[page_num].append((tbl_idx, rows))

    for page_num in sorted(pages.keys()):
        sheet_name = f"Page {page_num}"
        ws = wb.create_sheet(title=sheet_name)
        current_row = 1

        for tbl_idx, rows in pages[page_num]:
            if tbl_idx == 0:
                # plain text page — write as-is
                for line_row in rows:
                    for col_idx, cell_val in enumerate(line_row, start=1):
                        cell = ws.cell(row=current_row, column=col_idx, value=cell_val or "")
                        cell.alignment = Alignment(wrap_text=False)
                    current_row += 1
                current_row += 1
                continue

            # Write table header (first row)
            header = rows[0]
            for col_idx, cell_val in enumerate(header, start=1):
                cell = ws.cell(row=current_row, column=col_idx, value=str(cell_val or "").strip())
                cell.font  = HEADER_FONT
                cell.fill  = HEADER_FILL
                cell.border = BORDER
                cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            current_row += 1

            # Write data rows
            for row_offset, row in enumerate(rows[1:], start=0):
                fill = ALT_FILL if row_offset % 2 == 0 else NORM_FILL
                for col_idx, cell_val in enumerate(row, start=1):
                    val = str(cell_val or "").strip() if cell_val is not None else ""
                    # Try to coerce numbers
                    typed_val = _coerce(val)
                    cell = ws.cell(row=current_row, column=col_idx, value=typed_val)
                    cell.fill   = fill
                    cell.border = BORDER
                    cell.alignment = Alignment(wrap_text=True, vertical="top")
                    if isinstance(typed_val, (int, float)):
                        cell.number_format = '#,##0.00' if '.' in val else '#,##0'
                current_row += 1

            # Spacer between tables on the same page
            current_row += 1

        # Auto-fit column widths (capped at 60)
        for col in ws.columns:
            max_len = 0
            col_letter = get_column_letter(col[0].column)
            for cell in col:
                try:
                    max_len = max(max_len, len(str(cell.value or "")))
                except Exception:
                    pass
            ws.column_dimensions[col_letter].width = min(max(max_len + 2, 10), 60)

        # Freeze first row
        ws.freeze_panes = "A2"

    wb.save(output_path)
    print(f"\n[OK] Saved: {output_path}")


def _coerce(val: str):
    """Try to parse val as int or float; return original string on failure."""
    if not val:
        return val
    cleaned = val.replace(",", "").replace("$", "").strip()
    try:
        if "." in cleaned:
            return float(cleaned)
        return int(cleaned)
    except ValueError:
        return val


def calculate_totals(xlsx_path: Path):
    """
    Read an existing Excel workbook, sum column I (Proceeds) per security symbol
    across all sheets, and write/replace a 'Totals' sheet with the grouped results.

    Trade rows are identified by column A matching an IBKR account ID pattern (e.g. U10476244).
    Column B = Symbol, Column I = Proceeds.
    """
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter
    from collections import defaultdict
    import re

    if not xlsx_path.exists():
        print(f"[ERROR] Excel file not found: {xlsx_path}", file=sys.stderr)
        sys.exit(1)

    print(f"[INFO] Reading: {xlsx_path}")
    wb = openpyxl.load_workbook(xlsx_path)

    # Remove stale Totals sheet if present
    if "Totals" in wb.sheetnames:
        del wb["Totals"]

    acct_id_pattern = re.compile(r'^U\d+$')
    totals: dict[str, float] = defaultdict(float)
    trade_count: dict[str, int] = defaultdict(int)

    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        for row in ws.iter_rows(min_row=1, values_only=True):
            if len(row) < 9:
                continue
            acct_id = row[0]   # col A
            symbol  = row[1]   # col B
            proceeds = row[8]  # col I (0-indexed → index 8)

            if not (acct_id and symbol and isinstance(proceeds, (int, float))):
                continue
            if not acct_id_pattern.match(str(acct_id)):
                continue

            totals[str(symbol).strip()] += proceeds
            trade_count[str(symbol).strip()] += 1

    if not totals:
        print("[WARN] No trade rows found in column I. Nothing to total.")
        return

    # --- Style constants ---
    HEADER_FILL  = PatternFill("solid", fgColor="1F4E79")
    HEADER_FONT  = Font(bold=True, color="FFFFFF", size=11)
    TOTAL_FILL   = PatternFill("solid", fgColor="2E75B6")
    TOTAL_FONT   = Font(bold=True, color="FFFFFF", size=11)
    ALT_FILL     = PatternFill("solid", fgColor="D6E4F0")
    NORM_FILL    = PatternFill("solid", fgColor="FFFFFF")
    POS_FONT     = Font(color="1F6B2E")   # green for positive proceeds
    NEG_FONT     = Font(color="9C0006")   # red for negative proceeds
    THIN         = Side(style="thin", color="BFBFBF")
    BORDER       = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
    CENTER       = Alignment(horizontal="center", vertical="center")
    RIGHT        = Alignment(horizontal="right",  vertical="center")

    ws_totals = wb.create_sheet(title="Totals", index=0)

    # Header
    headers = ["Symbol", "Trade Count", "Total Proceeds"]
    for col_idx, h in enumerate(headers, start=1):
        cell = ws_totals.cell(row=1, column=col_idx, value=h)
        cell.font   = HEADER_FONT
        cell.fill   = HEADER_FILL
        cell.border = BORDER
        cell.alignment = CENTER

    # Data rows — sorted by symbol name
    grand_total = 0.0
    for row_offset, symbol in enumerate(sorted(totals.keys()), start=0):
        r = row_offset + 2
        fill = ALT_FILL if row_offset % 2 == 0 else NORM_FILL
        proceeds_val = totals[symbol]
        grand_total += proceeds_val

        sym_cell = ws_totals.cell(row=r, column=1, value=symbol)
        sym_cell.fill = fill; sym_cell.border = BORDER; sym_cell.alignment = CENTER

        cnt_cell = ws_totals.cell(row=r, column=2, value=trade_count[symbol])
        cnt_cell.fill = fill; cnt_cell.border = BORDER; cnt_cell.alignment = CENTER
        cnt_cell.number_format = '#,##0'

        proc_cell = ws_totals.cell(row=r, column=3, value=round(proceeds_val, 2))
        proc_cell.fill = fill; proc_cell.border = BORDER; proc_cell.alignment = RIGHT
        proc_cell.number_format = '#,##0.00'
        proc_cell.font = POS_FONT if proceeds_val >= 0 else NEG_FONT

    # Grand total row
    total_row = len(totals) + 2
    ws_totals.cell(row=total_row, column=1, value="GRAND TOTAL").font = TOTAL_FONT
    ws_totals.cell(row=total_row, column=1).fill   = TOTAL_FILL
    ws_totals.cell(row=total_row, column=1).border = BORDER
    ws_totals.cell(row=total_row, column=1).alignment = CENTER

    tc_cell = ws_totals.cell(row=total_row, column=2, value=sum(trade_count.values()))
    tc_cell.font = TOTAL_FONT; tc_cell.fill = TOTAL_FILL; tc_cell.border = BORDER
    tc_cell.alignment = CENTER; tc_cell.number_format = '#,##0'

    gt_cell = ws_totals.cell(row=total_row, column=3, value=round(grand_total, 2))
    gt_cell.font = TOTAL_FONT; gt_cell.fill = TOTAL_FILL; gt_cell.border = BORDER
    gt_cell.alignment = RIGHT; gt_cell.number_format = '#,##0.00'

    # Column widths
    ws_totals.column_dimensions["A"].width = 18
    ws_totals.column_dimensions["B"].width = 14
    ws_totals.column_dimensions["C"].width = 20
    ws_totals.freeze_panes = "A2"

    wb.save(xlsx_path)
    print(f"[OK] Totals sheet written — {len(totals)} symbols, grand total Proceeds = {grand_total:,.2f}")
    print(f"[OK] Saved: {xlsx_path}")


def main():
    parser = argparse.ArgumentParser(description="Convert PDF statement to Excel, preserving tables.")
    parser.add_argument("input_pdf", help="Path to input PDF file")
    parser.add_argument("output_xlsx", nargs="?", help="Path to output Excel file (default: same name as PDF)")
    parser.add_argument("--total", metavar="XLSX", nargs="?", const="__derive__",
                        help="Calculate per-symbol Proceeds totals from an existing Excel file "
                             "(defaults to the xlsx derived from input_pdf)")
    args = parser.parse_args()

    pdf_path = Path(args.input_pdf).resolve()
    xlsx_path = Path(args.output_xlsx).resolve() if args.output_xlsx else pdf_path.with_suffix(".xlsx")

    if args.total is not None:
        # --total mode: use explicitly supplied path or fall back to derived xlsx name
        target = xlsx_path if args.total == "__derive__" else Path(args.total).resolve()
        calculate_totals(target)
        return

    if not pdf_path.exists():
        print(f"[ERROR] File not found: {pdf_path}", file=sys.stderr)
        sys.exit(1)

    print(f"[INFO] Input:  {pdf_path}")
    print(f"[INFO] Output: {xlsx_path}")

    tables = extract_tables_pdfplumber(pdf_path)
    if not tables:
        print("[WARN] No content extracted from PDF.")
        sys.exit(1)

    write_excel(tables, xlsx_path)


if __name__ == "__main__":
    main()
